import docx
from pathlib import Path

doc = docx.Document("lecture_materials/05_CSS_Advanced_Properties.docx")

keywords = ["background", "font", "dimension", "position", "text", "width", "height"]
matches = []

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if any(k in t.lower() for k in ["background properties", "font properties", "text properties", "dimension properties", "positioning properties", "position properties"]):
        matches.append((i, t))

print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx, text in matches:
    print(f"Line {idx}: {text}")
    for j in range(idx+1, min(idx+12, len(doc.paragraphs))):
        if doc.paragraphs[j].text.strip():
            print(f"   -> {doc.paragraphs[j].text.strip()}")
