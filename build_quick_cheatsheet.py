import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DOCX_OUT = BASE_DIR / "QUICK_DEFENSE_CHEATSHEET.docx"
HTML_OUT = BASE_DIR / "QUICK_DEFENSE_CHEATSHEET.html"

COLOR_NAVY = RGBColor(15, 23, 42)
COLOR_ORANGE = RGBColor(234, 88, 12)
COLOR_SLATE = RGBColor(51, 65, 85)
COLOR_GREEN = RGBColor(22, 163, 74)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_point(doc, num, title, question, answer):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="30" w:space="0" w:color="EA580C"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    
    # Title & Question
    r_title = p.add_run(f"{num}. {title}\n")
    r_title.bold = True
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = COLOR_NAVY
    
    r_q = p.add_run(f"If Asked: \"{question}\"\n")
    r_q.italic = True
    r_q.font.size = Pt(9.5)
    r_q.font.color.rgb = COLOR_ORANGE
    
    r_ans_label = p.add_run("Your 10-Second Answer: ")
    r_ans_label.bold = True
    r_ans_label.font.size = Pt(9.5)
    r_ans_label.font.color.rgb = COLOR_GREEN
    
    r_ans = p.add_run(f"\"{answer}\"")
    r_ans.font.size = Pt(9.5)
    r_ans.font.color.rgb = COLOR_SLATE
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def build_docx():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("⚡ 2-MINUTE CODE DEFENSE CHEAT SHEET")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = COLOR_NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Web Programming Project 1 • BoBoiBoy Universe Portal • Resident University")
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = COLOR_ORANGE
    
    add_point(
        doc,
        num=1,
        title="Master Frame Architecture (index.html)",
        question="How does this website work?",
        answer="This website uses a Master Frame architecture. The navbar is written once in index.html, and content loads inside the <iframe name=\"mainFrame\"> using target=\"mainFrame\" without refreshing the outer navbar."
    )
    
    add_point(
        doc,
        num=2,
        title="The 5 Interconnected Pages",
        question="What are the 5 pages and what do they contain?",
        answer="We have 5 interconnected pages: home.html (lore & enlistment protocol), characters.html (hero cards & stat meters), elements.html (7 elements & matrix table), series.html (canon timeline & watch order), and contact.html (cadet form & student info)."
    )
    
    add_point(
        doc,
        num=3,
        title="Three Methods of Using CSS",
        question="Where are the 3 CSS methods implemented?",
        answer="1. External CSS in css/style.css for global styles, 2. Internal CSS in <style> tags in <head> for page components, and 3. Inline CSS in style=\"...\" attributes for micro-spacing and specific widths."
    )
    
    add_point(
        doc,
        num=4,
        title="Four Required CSS Property Categories",
        question="Name the 4 required CSS property categories!",
        answer="1. Background (color, image, position, repeat, size), 2. Text/Font (family, size, weight, line-height, color), 3. Dimensions (width, min-width, max-width, height, min-height), and 4. Positioning (position, top, bottom, left, right, z-index)."
    )
    
    add_point(
        doc,
        num=5,
        title="Form & Semantic HTML5 Validation (contact.html)",
        question="What form did you create and how is it validated?",
        answer="The Cadet Enlistment Form has all 7 required controls (text, password, textarea, radio, checkbox, select, submit). Validation is handled purely with HTML5 attributes like required and minlength=\"6\" with zero JavaScript."
    )
    
    add_point(
        doc,
        num=6,
        title="Clean Card Design & Light Ambient Shadow",
        question="Why does the design look clean and human-crafted rather than heavy AI?",
        answer="We used standard CSS Box Model spacing with a soft ambient shadow (box-shadow: 0 2px 4px rgba(0,0,0,0.06)) and crisp 1px borders, avoiding heavy artificial neon glow effects."
    )
    
    doc.save(DOCX_OUT)
    print(f"Generated {DOCX_OUT}")

def build_html():
    html = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>2-Minute Code Defense Cheat Sheet | BoBoiBoy Universe</title>
  <style>
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      background: #ffffff;
      color: #0f172a;
      line-height: 1.5;
      padding: 2rem 1rem;
      max-width: 800px;
      margin: 0 auto;
    }
    h1 {
      font-size: 1.6rem;
      color: #0f172a;
      margin-bottom: 0.25rem;
    }
    .sub {
      color: #ea580c;
      font-size: 0.9rem;
      font-weight: 600;
      margin-bottom: 1.5rem;
      border-bottom: 2px solid #ea580c;
      padding-bottom: 0.5rem;
    }
    .card {
      background: #f8fafc;
      border: 1px solid #e2e8f0;
      border-left: 4px solid #ea580c;
      border-radius: 6px;
      padding: 1rem 1.25rem;
      margin-bottom: 1rem;
    }
    .num {
      font-size: 1.05rem;
      font-weight: 800;
      color: #0f172a;
      margin-bottom: 0.35rem;
    }
    .q {
      font-size: 0.88rem;
      font-weight: 700;
      color: #ea580c;
      font-style: italic;
      margin-bottom: 0.35rem;
    }
    .ans {
      font-size: 0.92rem;
      color: #334155;
    }
    .ans strong {
      color: #16a34a;
    }
  </style>
</head>
<body>
  <h1>⚡ 2-Minute Code Defense Cheat Sheet</h1>
  <div class="sub">Web Programming Project 1 • BoBoiBoy Universe Portal • Resident University</div>

  <div class="card">
    <div class="num">1. Master Frame Architecture (index.html)</div>
    <div class="q">Dosen tanya: "Gimana cara kerja website ini?"</div>
    <div class="ans"><strong>Jawaban:</strong> "This website uses a <strong>Master Frame</strong> architecture. The navbar is written once in <code>index.html</code>, and content loads inside the <code>&lt;iframe name='mainFrame'&gt;</code> using <code>target='mainFrame'</code> without refreshing the outer navbar."</div>
  </div>

  <div class="card">
    <div class="num">2. The 5 Interconnected Pages</div>
    <div class="q">Dosen tanya: "Ada berapa halaman dan fungsinya apa saja?"</div>
    <div class="ans"><strong>Jawaban:</strong> "We have 5 interconnected pages: <code>home.html</code> (lore & enlistment protocol), <code>characters.html</code> (hero cards & stat meters), <code>elements.html</code> (7 elements & matrix table), <code>series.html</code> (canon timeline & watch order), and <code>contact.html</code> (cadet form & student info)."</div>
  </div>

  <div class="card">
    <div class="num">3. Three Methods of Using CSS</div>
    <div class="q">Dosen tanya: "Di mana 3 cara pakai CSS-nya?"</div>
    <div class="ans"><strong>Jawaban:</strong> "1. <strong>External CSS</strong> in <code>css/style.css</code> for global styles, 2. <strong>Internal CSS</strong> in <code>&lt;style&gt;</code> tags in <code>&lt;head&gt;</code> for page components, and 3. <strong>Inline CSS</strong> in <code>style='...'</code> attributes for micro-spacing and specific widths."</div>
  </div>

  <div class="card">
    <div class="num">4. Four Required CSS Property Categories</div>
    <div class="q">Dosen tanya: "Sebutkan 4 kategori CSS yang wajib di rubrik!"</div>
    <div class="ans"><strong>Jawaban:</strong> "1. <strong>Background</strong> (color, image, position, repeat, size), 2. <strong>Text/Font</strong> (family, size, weight, line-height, color), 3. <strong>Dimensions</strong> (width, min-width, max-width, height, min-height), and 4. <strong>Positioning</strong> (position, top, bottom, left, right, z-index)."</div>
  </div>

  <div class="card">
    <div class="num">5. Form & Semantic Validation (contact.html)</div>
    <div class="q">Dosen tanya: "Formulirnya ada apa saja dan validasinya pakai apa?"</div>
    <div class="ans"><strong>Jawaban:</strong> "The Cadet Enlistment Form has all <strong>7 required controls</strong> (text, password, textarea, radio, checkbox, select, submit). Validation is handled purely with HTML5 attributes like <code>required</code> and <code>minlength='6'</code> with zero JavaScript."</div>
  </div>

  <div class="card">
    <div class="num">6. Clean Card Design & Light Ambient Shadow</div>
    <div class="q">Dosen tanya: "Kenapa desain kotaknya rapi dan tidak berlebihan?"</div>
    <div class="ans"><strong>Jawaban:</strong> "We used standard <strong>CSS Box Model</strong> spacing with a soft ambient shadow (<code>box-shadow: 0 2px 4px rgba(0,0,0,0.06)</code>) and crisp 1px borders, avoiding heavy artificial neon glow effects."</div>
  </div>
</body>
</html>
"""
    with open(HTML_OUT, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"Generated {HTML_OUT}")

if __name__ == "__main__":
    build_docx()
    build_html()
